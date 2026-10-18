"""Shared building blocks for the pitch decks (build_pitch.py, build_pitch_v2.py).

The decks are built on the template docs/sources/Mamo_graf_VZP_v41.pptx: new_deck() opens it, drops its
slides and keeps its layouts ("TITUL", "OBSAH"), so the pitch keeps the W&P look (Poppins, navy palette,
logo, footers). The segment slide and the channel table exist for 7 segments (A–G) and 11 segments (A–G
split further); their numbers come from data_vzp/Rakhaton2026.csv.gz (see build_pitch.py).
"""
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
TEMPLATE = HERE.parent / "sources" / "Mamo_graf_VZP_v41.pptx"

NAVY, BLUE, LIGHT, CORAL, LILAC = "003B7A", "3D7CC9", "A9C6EA", "D27869", "7A78C3"
INK, BODY, MUTED, LINE, PALE, WHITE = "230F3C", "4A4458", "8A8597", "DEDCEC", "F7F6FB", "FFFFFF"
BADGE = {"Data VZP": NAVY, "Literatura": BLUE, "E-commerce": CORAL, "ÚZIS a ČSÚ": LILAC}
X0, XW = 0.85, 11.63          # content left edge and width (as in v41)

LAYOUT = {}
NOTES_BODY = None


def new_deck():
    """Open the template, drop its slides, remember its layouts and a notes body placeholder."""
    global NOTES_BODY
    prs = Presentation(TEMPLATE)
    # the template's notes master has no body placeholder, so new notes slides would have no text frame:
    # keep a copy of the notes body placeholder from an existing slide and reuse it
    NOTES_BODY = copy.deepcopy(next(sh for sh in prs.slides[0].notes_slide.placeholders
                                    if sh.placeholder_format.type == 2)._element)
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    LAYOUT.clear()
    LAYOUT.update({l.name: l for l in prs.slide_layouts})
    return prs


def nb(s):
    """Czech typography as in v41: non-breaking spaces before units, inside numbers, after one-letter words."""
    s = re.sub(r"(\d) (\d{3})", "\\1\u00a0\\2", s)
    s = re.sub(r" (%|p\. b\.|tis\.|mil\.|Kč|let|měs\.)", "\u00a0\\1", s)
    s = re.sub(r"(?<![\w])([vkszouaiVKSZOUAI]) ", "\\1\u00a0", s)
    return s.replace("p. b.", "p.\u00a0b.")


def rgb(h):
    return RGBColor.from_string(h)


def plain(shape):
    """Drop the theme style python-pptx adds to auto shapes (outline, shadow), so fill and line are ours."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def text(slide, x, y, w, h, content, size=13, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=0, bullets=False):
    """Text box; content is a string or a list of paragraphs; **bold** marks bold runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, ptext in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.18))))
            pPr.set("indent", str(-int(Inches(0.18))))
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
        for part in re.split(r"(\*\*[^*]+\*\*)", nb(ptext)):
            if not part:
                continue
            r = p.add_run()
            r.text = part[2:-2] if part.startswith("**") else part
            r.font.size = Pt(size)
            r.font.bold = bold or part.startswith("**")
            r.font.color.rgb = rgb(color)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.06):
    s = plain(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    return s


def badges(slide, kinds, x_right=X0 + XW, y=0.5, x_left=None):
    """Source badges (Data VZP / Literatura / E-commerce / ÚZIS a ČSÚ): right-aligned on the kicker row, or from x_left."""
    widths = [0.25 + 0.085 * len(k) for k in kinds]
    x = x_left if x_left is not None else x_right - sum(widths) - 0.08 * (len(kinds) - 1)
    for k, w in zip(kinds, widths):
        b = box(slide, x, y, w, 0.27, fill=BADGE[k], line=None, radius=0.5)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = k
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = rgb(WHITE)
        x += w + 0.08


def stat(slide, x, y, w, h, value, desc, color=NAVY):
    box(slide, x, y, w, h)
    text(slide, x + 0.22, y + 0.12, w - 0.44, 0.55, value, size=22, color=color, bold=True)
    text(slide, x + 0.22, y + 0.66, w - 0.44, h - 0.78, desc, size=10.5, color=BODY)


def source(slide, s):
    text(slide, 2.30, 6.88, 9.30, 0.30, s, size=9, color=MUTED)


def table(slide, x, y, w, col_w, rows, row_h=0.36, size=10.5, bold_rows=(), highlight_col=None):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = gf.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(r):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(NAVY if i == 0 else (PALE if i % 2 == 0 else WHITE))
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or (highlight_col is not None and j == 1) else PP_ALIGN.RIGHT
            for part in re.split(r"(\*\*[^*]+\*\*)", nb(str(val))):
                if not part:
                    continue
                run = p.add_run()
                run.text = part[2:-2] if part.startswith("**") else part
                run.font.size = Pt(size)
                run.font.bold = i == 0 or i in bold_rows or part.startswith("**")
                run.font.color.rgb = rgb(WHITE if i == 0 else INK)
    return gf


def arrow(slide, x, y, w=0.3, color=LIGHT):
    a = plain(slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.3)))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(color)
    a.line.fill.background()
    return a


def set_notes(slide, note):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(copy.deepcopy(NOTES_BODY))
    tf = ns.notes_text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    tf.text = note


def content_slide(prs, kicker, title, intro=None):
    s = prs.slides.add_slide(LAYOUT["OBSAH"])
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 101:
            ph.text_frame.text = kicker
        elif ph.placeholder_format.idx == 102:
            ph.text_frame.text = nb(title)
    if intro:
        text(s, X0, 1.62, XW, 0.5, intro, size=13, color=BODY)
    return s


def segments_slide(prs, seg):
    """Where the segments come from: the questions over VZP data and the 2024 table (7 or 11 segments)."""
    if seg == 7:
        s = content_slide(prs, "ODKUD SEGMENTY", "Tři otázky nad daty VZP rozdělí ženy do sedmi skupin",
                          "Pravidla stojí jen na polích, která VZP poskytla. Pořadí otázek je priorita: žena patří do první skupiny, "
                          "jejíž podmínka platí.")
        badges(s, ["Data VZP"])
        qs = [("1 · Historie screeningu", ["**A** screening před méně než 24 měs.", "**B** první pozvánka", "**C** screening v posledních 5 letech"]),
              ("2 · Prevence", ["**D** preventivní gynekologie", "**E** jiná prevence: zubař, praktik, jiný screening"]),
              ("3 · Kontakt se zdravotnictvím", ["**F** aspoň jedna návštěva lékaře za rok", "**G** žádný kontakt"])]
        y = 2.25
        for q, items in qs:
            h = 0.55 + 0.3 * len(items)
            box(s, X0, y, 4.35, h, fill=PALE)
            text(s, X0 + 0.2, y + 0.14, 4.0, 0.3, q, size=12, color=NAVY, bold=True)
            text(s, X0 + 0.2, y + 0.5, 4.0, h - 0.55, items, size=11, color=INK, space_after=3)
            y += h + 0.2
        table(s, 5.45, 2.25, 7.03, [3.13, 1.3, 1.3, 1.3], [
            ["Segment", "Ženy 2024", "Účast 12 měs.", "Nepřišlo"],
            ["A nedávno vyšetřená", "7 600", "2 %", "7 418"],
            ["B první pozvánka", "37 460", "55 %", "16 820"],
            ["C vypadlá z rytmu", "28 577", "33 %", "19 111"],
            ["D bez screeningu, gynekolog", "51 047", "21 %", "40 573"],
            ["E bez screeningu, jiná prevence", "89 351", "8 %", "82 559"],
            ["F jen léčebná péče", "28 368", "5 %", "26 977"],
            ["G mimo péči", "14 063", "2 %", "13 743"],
            ["Celkem", "256 466", "19 %", "207 201"],
        ], row_h=0.34, bold_rows=(8,))
        box(s, 5.45, 5.5, 7.03, 0.95, fill=WHITE)
        text(s, 5.65, 5.6, 6.65, 0.8, [
            "**Ověřeno na novém roce:** pravidla z let 2020–2023, test na 2024. Samotné segmenty seřadí ženy s AUC 0,79, "
            "celý vysvětlitelný model 0,83.",
            "Segmenty jsou stabilní pět let; jemnější verze s 11 segmenty slouží k cílení."], size=10.5, color=BODY, space_after=3)
        source(s, "Zdroj: data VZP pro Rakathon 2026, pozvánky 2020–2024; naše analýza")
        return s

    s = content_slide(prs, "ODKUD SEGMENTY", "Šest otázek nad daty VZP rozdělí ženy do jedenácti skupin",
                      "Pravidla stojí jen na polích, která VZP poskytla. Otázky jdou shora dolů; žena patří do první skupiny, "
                      "na kterou odpoví „ano“.")
    badges(s, ["Data VZP"])

    def chip_color(u):            # účast 2024: vysoká modře, střední světle, nízká korálově
        return (NAVY, WHITE) if u >= 0.3 else ((LIGHT, INK) if u >= 0.1 else (CORAL, WHITE))
    rows = [("Screening před < 24 měs.?", [("A", "dopis se minul", 0.024)]),
            ("První pozvánka?", [("B1", "< 6 kontaktů za rok", 0.357), ("B2", "v kontaktu", 0.596)]),
            ("Screening v posledních 5 letech?", [("C1", "≤ 48 měs.", 0.388), ("C2", "4–5+ let", 0.257)]),
            ("Preventivní gynekologie?", [("D1", "≤ 6 pozvánek", 0.271), ("D2", "7+ pozvánek", 0.103)]),
            ("Jiná prevence (zubař, praktik)?", [("E1", "≤ 6 pozvánek", 0.127), ("E2", "7+ pozvánek", 0.051)]),
            ("≥ 6 kontaktů se zdravotnictvím?", [("F", "ano: v kontaktu", 0.052), ("G", "ne: odpojená", 0.031)])]
    rh, rg, y = 0.56, 0.17, 2.25
    for i, (q, chips) in enumerate(rows):
        box(s, X0, y, 2.55, rh, fill=PALE)
        text(s, X0 + 0.12, y, 2.35, rh, q, size=10, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, X0 + 2.6, y + rh / 2 - 0.22, 0.4, 0.2, "ano", size=8.5, color="1B7837", bold=True, align=PP_ALIGN.CENTER)
        arrow(s, X0 + 2.62, y + rh / 2 - 0.05, 0.34)
        cx = X0 + 3.02
        for code, lbl, u in chips:
            fill, fg = chip_color(u)
            box(s, cx, y + 0.06, 1.18, rh - 0.12, fill=fill, line=None, radius=0.15)
            text(s, cx + 0.06, y + 0.06, 1.06, rh - 0.12, [f"**{code}** · {lbl}"], size=8.5, color=fg,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cx += 1.24
        if i < len(rows) - 1:
            ln = plain(s.shapes.add_connector(1, Inches(X0 + 0.4), Inches(y + rh), Inches(X0 + 0.4), Inches(y + rh + rg)))
            ln.line.color.rgb = rgb(MUTED)
            ln.line.width = Pt(1)
            text(s, X0 + 0.47, y + rh + 0.005, 0.3, rg, "ne", size=8, color="B2182B", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        y += rh + rg
    text(s, X0, y + 0.02, 5.4, 0.25, "barva: účast 2024 ≥ 30 % / 10–30 % / < 10 %", size=8.5, color=MUTED)
    table(s, 6.4, 2.25, 6.08, [2.98, 1.0, 1.0, 1.1], [
        ["Segment", "Ženy 2024", "Účast", "Nepřišlo"],
        ["A dopis se minul", "7 600", "2 %", "7 418"],
        ["B1 1. pozvánka, málo kontaktů", "7 127", "36 %", "4 580"],
        ["B2 1. pozvánka, v kontaktu", "30 333", "60 %", "12 240"],
        ["C1 vypadlá 3–4 roky", "16 167", "39 %", "9 892"],
        ["C2 vypadlá 4–5+ let", "12 410", "26 %", "9 219"],
        ["D1 gynekolog, zvaná málokrát", "31 098", "27 %", "22 679"],
        ["D2 gynekolog, nereaguje", "19 949", "10 %", "17 894"],
        ["E1 jiná prevence, zvaná málokrát", "29 487", "13 %", "25 734"],
        ["E2 jiná prevence, nereaguje", "59 864", "5 %", "56 825"],
        ["F bez prevence, v kontaktu", "18 102", "5 %", "17 155"],
        ["G odpojená", "24 329", "3 %", "23 565"],
        ["Celkem", "256 466", "19 %", "207 201"],
    ], row_h=0.29, size=9.5, bold_rows=(12,))
    box(s, 6.4, 6.08, 6.08, 0.62, fill=WHITE)
    text(s, 6.55, 6.13, 5.8, 0.52,
         "**Ověřeno na novém roce:** pravidla z let 2020–2023, test na 2024. Segmenty seřadí ženy s AUC 0,81, "
         "celý vysvětlitelný model 0,83.", size=9.5, color=BODY)
    source(s, "Zdroj: data VZP pro Rakathon 2026, pozvánky 2020–2024; config/segments_vzp_v2.yaml; naše analýza")
    return s


SEGMENTS_NOTES = {
    7: ("35 s. Odkud segmenty? Ze tří otázek nad daty, která VZP poskytla. Zaprvé historie screeningu: nedávno vyšetřená, první "
        "pozvánka, nebo vypadla z rytmu. Zadruhé, chodí na jinou prevenci, ke gynekologovi nebo k zubaři a praktikovi? Zatřetí, je "
        "vůbec v kontaktu se zdravotnictvím? Tak vznikne sedm skupin s účastí od 2 do 55 procent. Pravidla jsme stanovili na letech "
        "2020 až 2023 a otestovali na roce 2024: segmenty seřadí ženy skoro stejně dobře jako celý model. "
        "(AUC segmentů 0,789 je přepočet z dat; AUC 0,83 logistické regrese je model týmu z prezentace v41.)"),
    11: ("35 s. Odkud segmenty? Z šesti otázek nad daty, která VZP poskytla, kladených shora dolů. Nejdřív historie screeningu: "
         "nedávno vyšetřená, první pozvánka, vypadla z rytmu. Pak prevence: chodí ke gynekologovi, nebo k zubaři a praktikovi? "
         "U nich rozhoduje, kolik pozvánek už ignorovala. Nakonec, je vůbec v kontaktu se zdravotnictvím? Vznikne jedenáct "
         "skupin s účastí od 2 do 60 procent. Pravidla jsme stanovili na letech 2020 až 2023 a otestovali na roce 2024: "
         "segmenty seřadí ženy skoro stejně dobře jako celý model. "
         "(AUC segmentů 0,807 z notebooku vzp_segmentation_final; AUC 0,83 logistické regrese je model týmu z prezentace v41.)"),
}


def channels_table(s, seg):
    """Segment → channel and intervention → effect range from the literature (7 or 11 rows)."""
    if seg == 7:
        return table(s, X0, 2.25, 7.65, [2.1, 3.95, 1.6], [
            ["Segment", "Kanál a opatření", "Účinek (p. b.)"],
            ["A nedávno vyšetřená", "nezvat, ověřit výkony těsně před odesláním", "úspora 7 600 dopisů"],
            ["B první pozvánka", "pevný termín v pozvánce + SMS připomínka", "**+10 až +19**"],
            ["C vypadlá z rytmu", "pevný termín dřív: ve 25.–36. měsíci místo po 3–5 letech", "**+10 až +19**"],
            ["D gynekolog", "seznam žen bez mamografie pro gynekologa, objednání v ordinaci", "+2 až +6"],
            ["E jiná prevence", "praktik jako nový kanál: seznam žen, upozornění při prohlídce", "+2 až +6"],
            ["F, G", "zatím ne: nízká účast a nejistý účinek", "–"],
        ], row_h=0.58, highlight_col=1)
    return table(s, X0, 2.25, 7.65, [2.25, 4.0, 1.4], [
        ["Segment", "Kanál a opatření", "Účinek (p. b.)"],
        ["A dopis se minul", "nezvat, ověřit výkony těsně před odesláním", "úspora 7 600"],
        ["B1 1. pozvánka, málo kontaktů", "pevný termín + telefonát navíc k dopisu", "**+10 až +19**"],
        ["B2 1. pozvánka, v kontaktu", "pevný termín v pozvánce + SMS", "**+10 až +19**"],
        ["C1, C2 vypadlé", "pevný termín dřív: ve 25.–36. měsíci", "**+10 až +19**"],
        ["D1 gynekolog, zvaná málokrát", "seznam žen bez mamografie pro gynekologa", "+2 až +6"],
        ["D2 gynekolog, nereaguje", "objednání přímo v ordinaci místo dopisu", "+2 až +6"],
        ["E1 jiná prevence, zvaná málokrát", "praktik a zubař: seznam žen, upozornění", "+2 až +6"],
        ["E2 jiná prevence, nereaguje", "praktik místo dopisu; test vynechání dopisu", "+2 až +6"],
        ["F bez prevence, v kontaktu", "upozornění pro ošetřující lékaře (EZ karta)", "+1 až +5"],
        ["G odpojená", "zatím ne: nejdřív ověřit kontakt", "–"],
    ], row_h=0.38, size=9.5, highlight_col=1)
